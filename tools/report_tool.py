"""日报生成工具：聚合业务数据 → 生成 Excel 日报 → 推送企业微信。

对应实习工作：ODBC + SQL 取数 → 生成 Excel → 企业微信机器人自动推送日报。

[升级] 销售额为 0 时记入 warnings, Excel 路径 + 关键指标全结构化。
"""
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from config import DATA_DIR, WECHAT_WEBHOOK
from tools.validators import ToolResult, validate_tool, file_exists


@validate_tool(
    schema={"output_path": file_exists, "total_sales": lambda v: None},
    required_data_keys=("output_path", "total_sales", "order_count", "date_label"),
)
def generate_daily_report(date: str | None = None) -> ToolResult:
    """生成当日运营日报（Excel），并推送企业微信。"""
    clean_path = os.path.join(DATA_DIR, "sales_clean.csv")
    if not os.path.exists(clean_path):
        return ToolResult(
            success=False,
            errors=["尚未清洗销售数据，请先运行 clean_sales_data()"],
            summary="❌ 尚未清洗销售数据，请先运行 clean_sales_data()",
        )

    df = pd.read_csv(clean_path)
    df.columns = [str(c).strip().lower() for c in df.columns]

    date_col = "date" if "date" in df.columns else None
    if date_col:
        df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
        if date is None:
            target = df[date_col].max()
        else:
            target = pd.to_datetime(date)
        day_df = df[df[date_col] == target]
        date_label = target.strftime("%Y-%m-%d") if pd.notna(target) else "全部"
    else:
        day_df = df
        date_label = date or "全部"

    total_sales = float(day_df.get("amount", pd.Series([0])).sum()) if "amount" in day_df else 0.0
    order_count = int(len(day_df))
    top_products = ""
    if "product" in day_df.columns:
        grp = day_df.groupby("product")["amount"].sum().sort_values(ascending=False).head(5)
        top_products = "\n".join(f"      {p}: ¥{v:,.0f}" for p, v in grp.items())

    out_name = f"daily_report_{datetime.now():%Y%m%d}.xlsx"
    out_path = os.path.join(DATA_DIR, out_name)
    _write_excel(out_path, date_label, total_sales, order_count, top_products)

    summary = (
        f"📊 运营日报（{date_label}）\n"
        f"  · 销售额：¥{total_sales:,.0f}\n"
        f"  · 订单数：{order_count}\n"
        f"  · 报表文件：{out_name}"
    )
    _push_report(summary)
    summary += f"\n✅ 日报已生成：{out_path}"

    warnings: list[str] = []
    if total_sales == 0:
        warnings.append("当日销售额为 0，请确认数据日期是否正确")

    return ToolResult(
        success=True,
        data={
            "output_path": out_path,
            "total_sales": total_sales,
            "order_count": order_count,
            "date_label": date_label,
        },
        warnings=warnings,
        summary=summary,
    )


def _write_excel(path: str, date_label: str, total: float, orders: int, top: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "日报"
    title_font = Font(size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor="2C3E50")
    ws["A1"] = f"电商运营日报（{date_label}）"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws.merge_cells("A1:B1")
    rows = [
        ("统计日期", date_label),
        ("总销售额(¥)", f"{total:,.0f}"),
        ("订单数", orders),
        ("Top 商品", top or "无"),
    ]
    r = 3
    for k, v in rows:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, str(v)).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 48
    wb.save(path)


def _push_report(text: str) -> None:
    log_path = os.path.join(DATA_DIR, "wechat_alerts.log")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(text + "\n---\n")
    if WECHAT_WEBHOOK:
        try:
            import requests
            requests.post(WECHAT_WEBHOOK, json={"msgtype": "text", "text": {"content": text}}, timeout=5)
        except Exception:
            pass
